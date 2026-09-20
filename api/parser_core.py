"""
Painel Comercial A&B — lógica de leitura/tratamento dos 3 ficheiros Excel.
Versão sem pandas (só openpyxl) — mais leve e com menos dependências
binárias, para correr de forma fiável num ambiente serverless.
Partilhada entre a função serverless (api/refresh.py) e qualquer uso manual.
"""
import re
from datetime import datetime
from openpyxl import load_workbook

EMP_RE = re.compile(r'^st[mo]\d{3}a\d+$', re.I)
ID_STORE_RE = re.compile(r'^st([mo])(\d{3})a\d+$', re.I)
PERIODS = ['LD', 'WTD', 'MTD', 'QTD', 'YTD']

ADDITIVE = {'net_sales', 'sales_opt', 'sales_sun', 'sales_cl', 'sales_oth', 'cp_volume',
            'opt_volume', 'vaas', 'vaas_volume', 'insurance_sales', 'vs_tgt_eur'}
VOL_KEY = {'com': 'cp_volume', 'acc': 'opt_volume'}
ASP_KEY = {'com': 'cp_asp', 'acc': 'opt_asp'}

COM_KEYS = {
    'net_sales': 'NET SALES | Commercial Sales', 'vs_py_pct': 'NET SALES | vs PY %',
    'sales_opt': 'CATEGORY SALES | Net Sales OPT', 'sales_sun': 'CATEGORY SALES | Net Sales SUN',
    'sales_cl': 'CATEGORY SALES | Net Sales CL', 'sales_oth': 'CATEGORY SALES | Net Sales OTH',
    'cp_volume': 'COMPLETE PAIR | Volume', 'cp_asp': 'COMPLETE PAIR | ASP',
    'mf_pct': 'OPTICAL KPI | MF %', 'photochromic_pct': 'OPTICAL KPI | Photochromic %',
    'blue_pct': 'OPTICAL KPI | Blue %', 'vaas': 'OPTICAL KPI | VaaS €',
    'vaas_volume': 'OPTICAL KPI | VaaS Volume', 'vaas_share_pct': 'OPTICAL KPI | VaaS Share %',
}
ACC_KEYS = {
    'net_sales': 'NET SALES | Accounting Sales', 'vs_py_pct': 'NET SALES | vs PY %',
    'vs_tgt_eur': 'NET SALES | vs TGT €', 'vs_tgt_pct': 'NET SALES | vs TGT %',
    'sales_opt': 'CATEGORY NET SALES | Sales OPT', 'sales_sun': 'CATEGORY NET SALES | Sales SUN',
    'sales_cl': 'CATEGORY NET SALES | Sales CL', 'sales_oth': 'CATEGORY NET SALES | Sales OTH',
    'opt_volume': 'TOTAL OPTICAL | Volume', 'opt_asp': 'TOTAL OPTICAL | ASP',
    'mf_pct': 'OPTICAL KPI | MF %', 'photochromic_pct': 'OPTICAL KPI | Photochromic %',
    'blue_pct': 'OPTICAL KPI | Blue %', 'vaas': 'OPTICAL KPI | VaaS Sales',
    'vaas_volume': 'OPTICAL KPI | VaaS Volume', 'vaas_share_pct': 'OPTICAL KPI | VaaS Share',
    'insurance_sales': 'OTHER | Insurance Sales',
}


def g(rec, key):
    v = rec.get(key)
    return None if v in ('X', '#DIV/0', None) else v


def _sheet_rows(src, sheet_name):
    """Devolve a folha como uma lista de listas (1 linha = 1 lista), tal como
    pandas .values faria, para o resto do código poder indexar por posição
    sem se preocupar com a API do openpyxl."""
    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb[sheet_name]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def parse_kpi_sheet(src, sheet_name):
    rows = _sheet_rows(src, sheet_name)
    r1 = rows[1][1]
    r2, r3 = rows[2], rows[3]
    ncols = max(len(r) for r in rows)
    colnames, last_group = {}, None
    for i in range(6, ncols):
        sub = r2[i] if i < len(r2) and isinstance(r2[i], str) else None
        metric = r3[i] if i < len(r3) and isinstance(r3[i], str) else None
        if sub:
            last_group = sub
        if metric:
            colnames[i] = f"{last_group} | {metric}" if last_group else metric

    employees, stores, current_store = [], [], None
    for idx in range(4, len(rows)):
        row = rows[idx]
        country = row[1] if len(row) > 1 else None
        lux_id = row[5] if len(row) > 5 else None
        if isinstance(country, str) and country == 'PT':
            current_store = {'global_id': row[3], 'store_name': row[4]}
            continue
        if isinstance(row[1] if len(row) > 1 else None, str) and ' - ' in row[1] and (lux_id is None):
            gid, name = row[1].split(' - ', 1)
            rec = {'global_id': gid.strip(), 'store_name': name.strip()}
            for i, cn in colnames.items():
                v = row[i] if i < len(row) else None
                rec[cn] = None if v is None else v
            stores.append(rec)
            continue
        if lux_id is not None and current_store:
            lux_str = str(lux_id).strip()
            if EMP_RE.match(lux_str):
                rec = {'employee_id': lux_str.lower(), **current_store}
                for i, cn in colnames.items():
                    v = row[i] if i < len(row) else None
                    rec[cn] = None if v is None else v
                employees.append(rec)
    return {'period_label': r1, 'employees': employees, 'stores': stores}


def home_store_from_id(emp_id):
    m = ID_STORE_RE.match(emp_id)
    return f'5-{m.group(1).upper()}{m.group(2)}' if m else None


def combine_period_rows(rows, keys, kind):
    mapped = [{k: g(r, v) for k, v in keys.items()} | {'global_id': r['global_id'], 'store_name': r['store_name']}
              for r in rows]
    by_store = {m['global_id']: {k: m[k] for k in keys} | {'store_name': m['store_name']} for m in mapped}
    if len(mapped) == 1:
        total = dict(mapped[0])
        del total['global_id']; del total['store_name']
        return total, by_store
    total = {}
    for k in keys:
        vals = [m[k] for m in mapped if m.get(k) is not None]
        total[k] = round(sum(vals), 2) if (k in ADDITIVE and vals) else None
    dominant = max(mapped, key=lambda m: abs(m.get('net_sales') or 0))
    for k in keys:
        if total.get(k) is None and k not in ADDITIVE:
            total[k] = dominant.get(k)
    vk, ak = VOL_KEY[kind], ASP_KEY[kind]
    if vk in keys and ak in keys and total.get(vk):
        tot_sales = total.get('sales_opt')
        total[ak] = round(tot_sales / total[vk], 2) if tot_sales is not None else dominant.get(ak)
    return total, by_store


def _sales_report_rows(src):
    """Lê a folha 'LD Sales' do relatório de vendas: cabeçalho na linha 7
    (índice 6), dados a seguir. Devolve uma lista de dicts {coluna: valor}."""
    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb['LD Sales']
    all_rows = [list(row) for row in ws.iter_rows(values_only=True)]
    headers = all_rows[6]
    out = []
    for row in all_rows[7:]:
        rec = {}
        for i, h in enumerate(headers):
            if h is None:
                continue
            rec[h] = row[i] if i < len(row) else None
        out.append(rec)
    return out


def parse_all(dash_src, wo_src, sr_src):
    """dash_src/wo_src/sr_src: caminho de ficheiro OU objeto tipo-ficheiro
    (ex: io.BytesIO) para o dashboard, work-orders e sales-report.
    Devolve {'meta': {...}, 'stores': {...}, 'employees': {...}}."""
    kpi_data = {p: {k: parse_kpi_sheet(dash_src, f'{p}-{k}') for k in ['COM', 'ACC']} for p in PERIODS}

    snapshot_date = datetime.today().strftime('%d/%m/%Y')
    m = re.search(r'Date:\s*(\d{2}/\d{2}/\d{2})', str(kpi_data['LD']['COM']['period_label']))
    if m:
        d, mo, y = m.group(1).split('/')
        snapshot_date = f'{d}/{mo}/20{y}'

    sr_rows = _sales_report_rows(sr_src)
    txn_summary = {}
    by_emp = {}
    for rec in sr_rows:
        eid = rec.get('Employee ID')
        if not isinstance(eid, str) or not EMP_RE.match(eid.strip()):
            continue
        eid = eid.strip().lower()
        by_emp.setdefault(eid, []).append(rec)

    for eid, recs in by_emp.items():
        total_sales = sum((r.get('Net Sales') or 0) for r in recs)
        receipts = {r.get('Receipt Number') for r in recs if r.get('Receipt Number') is not None}
        prod_totals = {}
        for r in recs:
            if r.get('Sales Detail') != 'SALES':
                continue
            ns = r.get('Net Sales') or 0
            if ns <= 0:
                continue
            desc = r.get('SAP Article Desc')
            prod_totals[desc] = prod_totals.get(desc, 0) + ns
        top_products = dict(sorted(prod_totals.items(), key=lambda kv: kv[1], reverse=True)[:8])
        top_products = {k: round(v, 2) for k, v in top_products.items()}
        txn_summary[eid] = {
            'net_sales_today': round(float(total_sales), 2),
            'transactions': len(receipts),
            'top_products': top_products,
        }

    wo_rows = _sheet_rows(wo_src, 'RECAP BY STORE AND CLUSTER')
    aging, i = {}, 0
    while i < len(wo_rows):
        row = wo_rows[i]
        if len(row) > 1 and row[1] == 'Global Store ID':
            cluster_label, j = None, i + 1
            while j < len(wo_rows) and len(wo_rows[j]) > 1 and wo_rows[j][1] is not None and wo_rows[j][1] != 'Global Store ID':
                r = wo_rows[j]
                gid = r[1] if len(r) > 1 else None
                cluster = r[3] if len(r) > 3 else None
                amt = r[4] if len(r) > 4 else None
                if cluster is not None:
                    cluster_label = cluster
                if gid is not None and amt is not None:
                    aging.setdefault(gid, {})[cluster_label] = float(amt)
                j += 1
            i = j
        else:
            i += 1

    store_names = {}
    for p in PERIODS:
        for s in kpi_data[p]['ACC']['stores']:
            store_names[s['global_id']] = s['store_name']

    stores_out = {}
    for gid, name in store_names.items():
        periods_out = {}
        for p in PERIODS:
            com_list = [s for s in kpi_data[p]['COM']['stores'] if s['global_id'] == gid]
            acc_list = [s for s in kpi_data[p]['ACC']['stores'] if s['global_id'] == gid]
            pdata = {}
            if com_list:
                pdata['com'] = {k: g(com_list[0], v) for k, v in COM_KEYS.items()}
            if acc_list:
                pdata['acc'] = {k: g(acc_list[0], v) for k, v in ACC_KEYS.items()}
            periods_out[p] = pdata
        stores_out[gid] = {'name': name, 'aging': aging.get(gid, {}), 'periods': periods_out}

    emp_ids = {e['employee_id'] for p in PERIODS for e in kpi_data[p]['COM']['employees']}
    employees_out = {}
    for emp in sorted(emp_ids):
        home_gid = home_store_from_id(emp)
        periods_out = {}
        for p in PERIODS:
            com_rows = [e for e in kpi_data[p]['COM']['employees'] if e['employee_id'] == emp]
            acc_rows = [e for e in kpi_data[p]['ACC']['employees'] if e['employee_id'] == emp]
            pdata, com_by_store, acc_by_store = {}, {}, {}
            if com_rows:
                pdata['com'], com_by_store = combine_period_rows(com_rows, COM_KEYS, 'com')
            if acc_rows:
                pdata['acc'], acc_by_store = combine_period_rows(acc_rows, ACC_KEYS, 'acc')
            by_store = {}
            for gid in set(com_by_store) | set(acc_by_store):
                c, a = com_by_store.get(gid, {}), acc_by_store.get(gid, {})
                if not (c.get('net_sales') or a.get('net_sales')):
                    continue
                by_store[gid] = {
                    'store_name': c.get('store_name') or a.get('store_name') or store_names.get(gid, gid),
                    'com': {k: c[k] for k in COM_KEYS} if c else None,
                    'acc': {k: a[k] for k in ACC_KEYS} if a else None,
                }
            pdata['by_store'] = by_store
            periods_out[p] = pdata
        employees_out[emp] = {
            'store_gid': home_gid,
            'periods': periods_out,
            'today': txn_summary.get(emp),
        }

    return {
        'meta': {'snapshot_date': snapshot_date},
        'stores': stores_out,
        'employees': employees_out,
    }
